"""Bảng thống kê tự sinh — đúng hợp đồng, đếm đúng, và nhãn cảnh báo không rơi mất."""

from __future__ import annotations

import io

from openpyxl import load_workbook

from design_compute.contracts import validate
from design_compute.geometry.norms import load_construction_norms
from design_compute.schedules import DISCLAIMER, build_schedules, schedules_to_xlsx
from tests import test_geometry as tg

REF = f"sha256:{'a' * 64}"


def _plan():
    return tg.TestDxfExport._plan()


class TestSchedules:
    def test_matches_contract_and_counts_every_opening(self) -> None:
        plan = _plan()
        result = build_schedules(plan, floorplan_ref=REF, norms=load_construction_norms())
        validate("schedules", result)

        openings = plan["levels"][0]["openings"]
        doors = sum(r["count"] for r in result["doors"])
        windows = sum(r["count"] for r in result["windows"])
        assert doors == len([o for o in openings if o["kind"] == "door"])
        assert windows == len([o for o in openings if o["kind"] == "window"])
        assert result["disclaimer"] == DISCLAIMER

    def test_same_size_shares_a_code_and_codes_go_from_wide_to_narrow(self) -> None:
        result = build_schedules(_plan(), floorplan_ref=REF, norms=load_construction_norms())
        widths = [r["w_m"] for r in result["doors"] if not r["code"].startswith("DW")]
        assert widths == sorted(widths, reverse=True)
        assert len({r["code"] for r in result["doors"]}) == len(result["doors"])

    def test_areas_add_up_to_the_floor(self) -> None:
        plan = _plan()
        result = build_schedules(plan, floorplan_ref=REF, norms=load_construction_norms())
        assert round(sum(r["area_m2"] for r in result["areas"]), 1) == round(
            sum(r["area_m2"] for r in plan["levels"][0]["rooms"]), 1
        )
        floor_row = next(m for m in result["materials"] if m["code"] == "SAN")
        assert floor_row["area_m2"] == round(sum(r["area_m2"] for r in result["areas"]), 1)

    def test_materials_come_from_construction_norms_not_code(self) -> None:
        norms = load_construction_norms()
        result = build_schedules(_plan(), floorplan_ref=REF, norms=norms)
        assert all(r["material"] in set(norms.schedule_materials.values()) for r in result["doors"])

    def test_xlsx_has_three_sheets_with_the_disclaimer_on_top(self) -> None:
        result = build_schedules(_plan(), floorplan_ref=REF, norms=load_construction_norms())
        data = schedules_to_xlsx(result, labels={"living": "Phòng khách"}, title="NVO-028")
        wb = load_workbook(io.BytesIO(data))
        assert wb.sheetnames == ["Thống kê cửa", "Thống kê diện tích", "Khối lượng sơ bộ"]
        for ws in wb.worksheets:
            assert ws["A2"].value == DISCLAIMER
        names = [row[1] for row in wb["Thống kê diện tích"].iter_rows(min_row=5, values_only=True)]
        assert "Phòng khách" in names
